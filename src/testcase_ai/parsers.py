"""把需求和知识源文件解析成平台可处理的纯文本与标准知识块。"""

from __future__ import annotations

import hashlib
import re
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook

from testcase_ai.contracts import ExcelMapping, SourceType

SUPPORTED_KNOWLEDGE_SUFFIXES = {".md", ".markdown", ".txt", ".docx", ".xlsx"}


def sha256_text(value: str) -> str:
    """返回文本的稳定 SHA-256，用于 logical ID、内容哈希和版本指纹。"""

    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def normalize_text(value: Any) -> str:
    """把 Excel 等外部输入归一化为去除多余横向空白的字符串。"""

    if value is None:
        return ""
    return re.sub(r"[ \t\r\f\v]+", " ", str(value)).strip()


@dataclass(slots=True)
class ParsedChunk:
    """解析阶段的中间知识块，稍后会分别写入 MySQL 和向量库。"""

    logical_id: str
    content: str
    source_path: str
    source_type: SourceType
    section: str | None = None
    module: str | None = None
    parent_logical_id: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def content_hash(self) -> str:
        """根据块正文实时计算内容哈希，便于审计与变更识别。"""

        return sha256_text(self.content)


def parse_requirement_file(path: Path) -> tuple[str, str]:
    """解析 Markdown、TXT 或 DOCX 需求，返回推断出的标题和完整正文。"""

    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        content = path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        document = Document(path)
        content = "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
    else:
        raise ValueError(f"unsupported requirement format: {suffix}")
    title = next((line.lstrip("# ").strip() for line in content.splitlines() if line.strip()), path.stem)
    return title or path.stem, content.strip()


def split_markdown_sections(text: str) -> list[tuple[str, str]]:
    """按 Markdown 标题拆分章节，并保留标题行作为章节上下文。"""

    sections: list[tuple[str, list[str]]] = []
    current_heading = "正文"
    current_lines: list[str] = []
    for raw_line in text.splitlines():
        heading = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*$", raw_line)
        if heading:
            if current_lines:
                sections.append((current_heading, current_lines))
            current_heading = heading.group(1).strip()
            current_lines = [raw_line]
        else:
            current_lines.append(raw_line)
    if current_lines:
        sections.append((current_heading, current_lines))
    return [(heading, "\n".join(lines).strip()) for heading, lines in sections if "\n".join(lines).strip()]


def _window_text(text: str, chunk_size: int, overlap: int) -> Iterable[str]:
    """按长度滑窗切块，优先在换行或句号处分隔，并保留指定重叠文本。"""

    if len(text) <= chunk_size:
        yield text
        return
    start = 0
    while start < len(text):
        end = min(len(text), start + chunk_size)
        if end < len(text):
            split_at = max(text.rfind("\n", start, end), text.rfind("。", start, end))
            if split_at > start + chunk_size // 2:
                end = split_at + 1
        yield text[start:end].strip()
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)


def parse_text_knowledge(
    path: Path,
    *,
    relative_path: str,
    source_type: SourceType,
    chunk_size: int,
    overlap: int,
) -> list[ParsedChunk]:
    """将文本类知识按章节和滑窗切块，并附加来源、章节及父块信息。"""

    if path.suffix.lower() == ".docx":
        document = Document(path)
        text = "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
        sections = [(path.stem, text)]
    else:
        text = path.read_text(encoding="utf-8")
        sections = (
            split_markdown_sections(text)
            if path.suffix.lower() in {".md", ".markdown"}
            else [(path.stem, text)]
        )

    chunks: list[ParsedChunk] = []
    for section_index, (section, section_text) in enumerate(sections):
        parent_id = sha256_text(f"{relative_path}:{section_index}:{section}")
        windows = list(_window_text(section_text, chunk_size, overlap))
        for window_index, content in enumerate(windows):
            logical_id = sha256_text(f"{parent_id}:{window_index}")
            chunks.append(
                ParsedChunk(
                    logical_id=logical_id,
                    parent_logical_id=parent_id if len(windows) > 1 else None,
                    content=content,
                    source_path=relative_path,
                    source_type=source_type,
                    section=section,
                    module=path.stem if source_type == SourceType.BUSINESS_RULE else None,
                    metadata={"section_index": section_index, "window_index": window_index},
                )
            )
    return chunks


def _header_positions(worksheet: Any, header_row: int) -> dict[str, int]:
    """建立 Excel 表头名称到列号的映射，忽略空表头和重复表头。"""

    positions: dict[str, int] = {}
    for index, cell in enumerate(worksheet[header_row], start=1):
        header = normalize_text(cell.value)
        if header and header not in positions:
            positions[header] = index
    return positions


def _cell(worksheet: Any, row: int, positions: dict[str, int], header: str) -> str:
    """按表头安全读取指定行单元格；缺少可选列时返回空字符串。"""

    column = positions.get(header)
    return normalize_text(worksheet.cell(row=row, column=column).value) if column else ""


def parse_excel_cases(
    path: Path,
    *,
    relative_path: str,
    source_type: SourceType,
    mapping: ExcelMapping,
) -> list[ParsedChunk]:
    """把正式用例 Excel 的每一条非空用例转换成一个独立知识块。"""

    workbook = load_workbook(path, read_only=False, data_only=True)
    if mapping.sheet not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"{relative_path} missing sheet: {mapping.sheet}")
    worksheet = workbook[mapping.sheet]
    positions = _header_positions(worksheet, mapping.header_row)
    required = {"用例编号", "功能模块", "测试点", "测试步骤", "预期结果"}
    missing = sorted(required - positions.keys())
    if missing:
        workbook.close()
        raise ValueError(f"{relative_path} missing testcase columns: {', '.join(missing)}")

    chunks: list[ParsedChunk] = []
    for row_index in range(mapping.data_start_row, worksheet.max_row + 1):
        case_id = _cell(worksheet, row_index, positions, "用例编号")
        module = _cell(worksheet, row_index, positions, "功能模块")
        test_point = _cell(worksheet, row_index, positions, "测试点")
        expected = _cell(worksheet, row_index, positions, "预期结果")
        if not any((case_id, module, test_point, expected)):
            continue
        fields = {
            "用例编号": case_id,
            "适用版本": _cell(worksheet, row_index, positions, "适用版本"),
            "功能模块": module,
            "测试路径": _cell(worksheet, row_index, positions, "测试路径"),
            "测试点": test_point,
            "前置条件": _cell(worksheet, row_index, positions, "前置条件"),
            "测试步骤": _cell(worksheet, row_index, positions, "测试步骤"),
            "测试数据": _cell(worksheet, row_index, positions, "测试数据"),
            "预期结果": expected,
        }
        content = "\n".join(f"{key}: {value}" for key, value in fields.items() if value)
        identity = case_id or sha256_text(content)
        chunks.append(
            ParsedChunk(
                logical_id=sha256_text(f"{relative_path}:{identity}"),
                content=content,
                source_path=relative_path,
                source_type=source_type,
                section=mapping.sheet,
                module=module or path.stem,
                metadata={"case_id": case_id, "row": row_index, "fields": fields},
            )
        )
    workbook.close()
    return chunks


def parse_knowledge_file(
    path: Path,
    *,
    relative_path: str,
    source_type: SourceType,
    excel_mapping: ExcelMapping,
    chunk_size: int,
    overlap: int,
) -> list[ParsedChunk]:
    """根据文件后缀分派到 Excel 或文本解析器，不支持的类型返回空列表。"""

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_KNOWLEDGE_SUFFIXES:
        return []
    if suffix == ".xlsx":
        return parse_excel_cases(
            path,
            relative_path=relative_path,
            source_type=source_type,
            mapping=excel_mapping,
        )
    return parse_text_knowledge(
        path,
        relative_path=relative_path,
        source_type=source_type,
        chunk_size=chunk_size,
        overlap=overlap,
    )
