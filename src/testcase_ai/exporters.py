"""把标准生成结果导出为可审计 JSON 和兼容现有模板的 Excel 草稿。"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from testcase_ai.contracts import GenerationResultV1, TestCaseV1

HEADERS = [
    "用例编号",
    "适用版本",
    "功能模块",
    "测试路径",
    "测试点",
    "前置条件",
    "测试步骤",
    None,
    None,
    "测试数据",
    "预期结果",
    "实际结果",
    "通过 / 未通过",
    "测试日期",
]


def _joined(values: list[str]) -> str:
    """清理并用中文分号拼接多值字段，便于写入单个 Excel 单元格。"""

    return "；".join(item.strip() for item in values if item.strip())


def testcase_row(case: TestCaseV1) -> list[str | None]:
    """按照正式 Excel 的固定十四列顺序把一条标准用例转换为行数据。"""

    return [
        case.case_id,
        ",".join(case.applicable_versions),
        case.module,
        case.test_path,
        case.test_point,
        _joined(case.preconditions),
        _joined(case.steps),
        None,
        None,
        _joined(case.test_data),
        case.expected_result,
        case.actual_result,
        case.execution_status,
        case.test_date.isoformat() if case.test_date else None,
    ]


class ArtifactExporter:
    """负责生成任务输出目录及 JSON/XLSX 两种草稿文件。"""

    def __init__(self, outputs_root: Path) -> None:
        """固定输出根目录，所有产物按 project_id/generation_id 隔离。"""

        self.outputs_root = outputs_root.resolve()

    def export(self, result: GenerationResultV1, formats: list[str]) -> dict[str, str]:
        """按请求格式写出产物，并返回格式名到绝对文件路径的映射。"""

        target = self.outputs_root / result.project_id / result.generation_id
        target.mkdir(parents=True, exist_ok=True)
        artifacts = {
            format_name: str(target / f"testcases.{format_name}")
            for format_name in formats
            if format_name in {"json", "xlsx"}
        }
        if "json" in formats:
            path = Path(artifacts["json"])
            path.write_text(
                json.dumps(
                    result.model_copy(update={"artifacts": artifacts}).model_dump(mode="json"),
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
        if "xlsx" in formats:
            path = Path(artifacts["xlsx"])
            self._write_xlsx(result, path)
        return artifacts

    @staticmethod
    def _write_xlsx(result: GenerationResultV1, path: Path) -> None:
        """创建测试用例主表与来源映射表，并设置表头、合并单元格和可读样式。"""

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "全部用例"
        worksheet.merge_cells("A1:N1")
        worksheet["A1"] = f"{result.requirement.title} - 测试用例草稿"
        worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
        worksheet["A1"].fill = PatternFill("solid", fgColor="548235")
        worksheet["A1"].alignment = Alignment(horizontal="center")
        worksheet["A2"] = "项目"
        worksheet["B2"] = result.project_id
        worksheet["G2"] = "知识版本"
        worksheet["H2"] = result.knowledge_revision_id
        worksheet["A3"] = "测试摘要"
        worksheet["B3"] = result.requirement.title
        worksheet["G3"] = "状态"
        worksheet["H3"] = result.status.value

        for column, header in enumerate(HEADERS, start=1):
            if header:
                cell = worksheet.cell(row=4, column=column, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="C6E0B4")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.merge_cells("G4:I4")
        worksheet.freeze_panes = "A5"

        for row_index, case in enumerate(result.cases, start=5):
            values = testcase_row(case)
            for column, value in enumerate(values, start=1):
                if column in {8, 9}:
                    continue
                cell = worksheet.cell(row=row_index, column=column, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            worksheet.merge_cells(start_row=row_index, start_column=7, end_row=row_index, end_column=9)
            worksheet.row_dimensions[row_index].height = max(
                32, min(120, 16 * max(2, len(_joined(case.steps)) // 24))
            )

        widths = {1: 20, 2: 26, 3: 14, 4: 32, 5: 28, 6: 28, 7: 42, 10: 24, 11: 32, 12: 24, 13: 14, 14: 14}
        for column, width in widths.items():
            worksheet.column_dimensions[worksheet.cell(row=4, column=column).column_letter].width = width

        mapping = workbook.create_sheet("来源映射")
        mapping.append(["用例编号", "证据编号", "来源文件", "章节", "证据状态"])
        for case in result.cases:
            if not case.citations:
                mapping.append([case.case_id, "", "", "", case.evidence_status])
            for citation in case.citations:
                mapping.append(
                    [
                        case.case_id,
                        citation.evidence_id,
                        citation.source,
                        citation.section or "",
                        case.evidence_status,
                    ]
                )
        for cell in mapping[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="44546A")
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        mapping.freeze_panes = "A2"
        mapping.auto_filter.ref = mapping.dimensions
        for column, width in enumerate([20, 38, 60, 32, 20], start=1):
            mapping.column_dimensions[mapping.cell(row=1, column=column).column_letter].width = width
        workbook.save(path)
        workbook.close()
