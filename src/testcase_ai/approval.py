"""人工审批主链路：安全地把选中草稿写入正式 Excel，并触发知识重新索引。"""

from __future__ import annotations

import os
import re
import shutil
import tempfile
import uuid
from collections import defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from testcase_ai.contracts import ApprovalResultV1, GenerationStatus, TestCaseV1
from testcase_ai.database import ApprovalRecord, Database
from testcase_ai.exporters import testcase_row
from testcase_ai.generation import GenerationService
from testcase_ai.knowledge import KnowledgeService
from testcase_ai.manifest import resolve_project


class ApprovalService:
    """协调生成记录、正式资产写盘、备份、幂等记录与知识版本发布。"""

    def __init__(
        self,
        *,
        database: Database,
        generation: GenerationService,
        knowledge: KnowledgeService,
        projects_root: Path,
        outputs_root: Path,
    ) -> None:
        """注入审批依赖，并固定项目与输出根目录，防止相对路径漂移。"""

        self.database = database
        self.generation = generation
        self.knowledge = knowledge
        self.projects_root = projects_root.resolve()
        self.outputs_root = outputs_root.resolve()

    def approve(self, generation_id: str, selected_case_ids: list[str], *, confirm: bool) -> ApprovalResultV1:
        """显式确认后将选中用例原子导入正式工作簿。

        函数先检查 confirm、任务状态、用例归属和澄清状态，再按模块准备临时副本与备份。
        所有工作簿校验通过后才用 ``os.replace`` 替换正式文件；异常时从备份恢复。成功记录
        approval_id 后重新同步知识，使新增正式用例进入下一版检索索引。重复审批同一生成任务
        会直接返回第一次的记录，不会重复写入。
        """

        if not confirm:
            raise ValueError("formal import requires confirm=true")
        with self.database.session() as session:
            existing = session.query(ApprovalRecord).filter_by(generation_id=generation_id).one_or_none()
            if existing:
                return ApprovalResultV1(
                    approval_id=existing.id,
                    generation_id=generation_id,
                    project_id=existing.project_id,
                    imported_case_ids=existing.selected_case_ids,
                    target_files=existing.target_files,
                )

        result = self.generation.get(generation_id)
        if result.status not in {GenerationStatus.SUCCEEDED, GenerationStatus.NEEDS_CLARIFICATION}:
            raise ValueError("only a successful generation can be approved")
        selected_set = set(selected_case_ids)
        selected = [case for case in result.cases if case.case_id in selected_set]
        missing = sorted(selected_set - {case.case_id for case in selected})
        if missing:
            raise ValueError(f"generation does not contain case ids: {', '.join(missing)}")
        if any(case.evidence_status == "needs_clarification" for case in selected):
            raise ValueError("cases that need clarification cannot be approved")

        project_dir, manifest = resolve_project(self.projects_root, result.project_id)
        approved_root = project_dir / "knowledge" / "approved-cases" / "modules"
        grouped: dict[str, list[TestCaseV1]] = defaultdict(list)
        for case in selected:
            grouped[case.module].append(case)

        approval_id = str(uuid.uuid4())
        backup_root = self.outputs_root / result.project_id / "approvals" / approval_id / "backups"
        backup_root.mkdir(parents=True, exist_ok=True)
        prepared: list[tuple[Path, Path, Path]] = []
        imported_ids: list[str] = []
        try:
            for module, cases in grouped.items():
                safe_module = re.sub(r"[\\/:*?\"<>|]", "_", module)
                target = approved_root / f"{safe_module}.xlsx"
                if not target.is_file():
                    raise ValueError(f"approved module workbook not found: {target}")
                backup = backup_root / target.name
                shutil.copy2(target, backup)
                descriptor, temp_name = tempfile.mkstemp(
                    prefix=f".{target.stem}-", suffix=".xlsx", dir=target.parent
                )
                os.close(descriptor)
                temp_path = Path(temp_name)
                shutil.copy2(target, temp_path)
                workbook = load_workbook(temp_path)
                worksheet = workbook[manifest.excel.sheet]
                existing_ids = {
                    str(worksheet.cell(row=row, column=1).value or "").strip()
                    for row in range(manifest.excel.data_start_row, worksheet.max_row + 1)
                }
                counter = 1
                for case in cases:
                    formal_id = case.case_id
                    if formal_id.startswith("DRAFT-") or formal_id in existing_ids:
                        while True:
                            formal_id = f"AI-{datetime.now():%Y%m%d}-{counter:04d}"
                            counter += 1
                            if formal_id not in existing_ids:
                                break
                    row_case = case.model_copy(update={"case_id": formal_id})
                    worksheet.append(testcase_row(row_case))
                    row_index = worksheet.max_row
                    worksheet.merge_cells(
                        start_row=row_index, start_column=7, end_row=row_index, end_column=9
                    )
                    for cell in worksheet[row_index]:
                        alignment = copy(cell.alignment)
                        alignment.wrap_text = True
                        alignment.vertical = "top"
                        cell.alignment = alignment
                    existing_ids.add(formal_id)
                    imported_ids.append(formal_id)
                workbook.save(temp_path)
                workbook.close()
                check = load_workbook(temp_path, read_only=True, data_only=True)
                if manifest.excel.sheet not in check.sheetnames:
                    check.close()
                    raise ValueError(f"prepared workbook is invalid: {target}")
                check.close()
                prepared.append((target, temp_path, backup))

            for target, temp_path, _ in prepared:
                os.replace(temp_path, target)
        except Exception:
            for target, temp_path, backup in prepared:
                if temp_path.exists():
                    temp_path.unlink()
                if backup.exists():
                    shutil.copy2(backup, target)
            raise

        target_files = [str(target) for target, _, _ in prepared]
        with self.database.session() as session:
            session.add(
                ApprovalRecord(
                    id=approval_id,
                    generation_id=generation_id,
                    project_id=result.project_id,
                    selected_case_ids=imported_ids,
                    target_files=target_files,
                )
            )
        sync_result = self.knowledge.sync(project_dir)
        return ApprovalResultV1(
            approval_id=approval_id,
            generation_id=generation_id,
            project_id=result.project_id,
            imported_case_ids=imported_ids,
            target_files=target_files,
            knowledge_revision_id=sync_result.revision_id,
        )
